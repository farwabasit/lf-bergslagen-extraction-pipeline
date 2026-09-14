import io
import mimetypes

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_EXTRACTED_CHARS = 8000


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(v) for v in row if v is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _unsupported_binary_message(filename: str, kind: str) -> str:
    return (
        f"[{kind} attachment: {filename}]\n"
        f"The file was attached successfully, but this legacy {kind.lower()} format can't be "
        "read for text extraction. Please re-save it as .docx / .xlsx (or PDF) if you'd like "
        "Sara to read its contents."
    )


def extract_text(filename: str, data: bytes) -> str:
    lower_name = filename.lower()
    content_type, _ = mimetypes.guess_type(filename)

    if lower_name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif lower_name.endswith(".docx"):
        text = _extract_docx(data)
    elif lower_name.endswith(".xlsx"):
        text = _extract_xlsx(data)
    elif lower_name.endswith(".doc"):
        text = _unsupported_binary_message(filename, "Word document")
    elif lower_name.endswith(".xls"):
        text = _unsupported_binary_message(filename, "Excel spreadsheet")
    elif content_type and content_type.startswith("image/"):
        text = (
            f"[Image attachment: {filename}]\n"
            "The image was attached successfully, but text extraction from images is not enabled."
        )
    else:
        text = data.decode("utf-8", errors="ignore")
    return text.strip()[:MAX_EXTRACTED_CHARS]
